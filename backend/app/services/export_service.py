from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

from docx import Document
from docx.shared import Inches
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.business_plan import BusinessPlan
from app.models.finance_forecast import FinanceForecast
from app.models.financial_assumption import FinancialAssumption
from app.models.project import Project
from app.models.research_report import ResearchReport
from app.schemas.export import ExportFile, ExportResult
from app.finance.narrative import build_financial_narrative
from app.services.language_service import normalize_to_language

EXPORT_DIR = Path(__file__).resolve().parents[2] / "exports"


async def create_pdf_export(db: AsyncSession, project: Project, include_review: bool = False) -> ExportResult:
    package = await _load_export_package(db, project, include_review)
    file_name = _file_name(project.id, "pdf")
    file_path = _export_path(file_name)
    _write_pdf(file_path, package)
    return _to_result(file_name)


async def create_docx_export(db: AsyncSession, project: Project, include_review: bool = False) -> ExportResult:
    package = await _load_export_package(db, project, include_review)
    file_name = _file_name(project.id, "docx")
    file_path = _export_path(file_name)
    _write_docx(file_path, package)
    return _to_result(file_name)


def list_exports(project_id: str) -> list[ExportFile]:
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    files = sorted(
        EXPORT_DIR.glob(f"{project_id}_*.*"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    return [
        ExportFile(
            file_name=path.name,
            download_url=f"/exports/{path.name}",
            format=path.suffix.lstrip("."),
        )
        for path in files
        if path.suffix in {".pdf", ".docx"}
    ]


async def _load_export_package(db: AsyncSession, project: Project, include_review: bool = False) -> dict:
    forecast = await _latest(db, FinanceForecast, project.id)
    business_plan = await _latest(db, BusinessPlan, project.id)
    research = await _latest(db, ResearchReport, project.id)
    assumptions = await _latest(db, FinancialAssumption, project.id)

    missing = []
    if not forecast:
        missing.append("finance forecast")
    if not business_plan:
        missing.append("business plan")
    if not research:
        missing.append("research report")
    if missing:
        raise ValueError(f"Missing required data: {', '.join(missing)}")

    bp_content = _normalize_business_plan_content(business_plan.sections or {})
    forecast_data = forecast.summary or {}
    forecast_summary = forecast_data.get("summary", forecast_data)
    financial_narrative = build_financial_narrative(forecast)

    review_score = None
    if include_review:
        from app.services.review_service import latest_review_score

        review_score = await latest_review_score(db, project.id)

    return {
        "project_title": project.title,
        "idea_summary": project.idea_summary or "",
        "generated_date": datetime.now(timezone.utc).strftime("%Y-%m-%d"),
        "language": bp_content["language"],
        "template_type": bp_content["template_type"],
        "selected_sections": bp_content["selected_sections"],
        "business_plan_sections": bp_content["sections"],
        "market_size": research.market_size or "",
        "competitors": research.competitors or [],
        "industry_trends": research.industry_trends or [],
        "research_risks": research.risks or [],
        "citations": research.citations or [],
        "forecast_summary": forecast_summary,
        "forecast_data": forecast_data,
        "financial_narrative": financial_narrative,
        "financial_assumptions": assumptions.assumptions if assumptions else {},
        "financial_interpretation": forecast_data.get("interpretation", {}),
        "chart_files": forecast_data.get("chart_files", {}),
        "breakeven": forecast_summary.get("breakeven_month"),
        "funding_needed": forecast_summary.get("funding_needed"),
        "review_score": review_score,
    }


async def _latest(db: AsyncSession, model, project_id: str):
    result = await db.execute(
        select(model)
        .where(model.project_id == project_id)
        .order_by(model.created_at.desc())
        .limit(1)
    )
    return result.scalar_one_or_none()


def _file_name(project_id: str, extension: str) -> str:
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S")
    return f"{project_id}_{timestamp}.{extension}"


def _export_path(file_name: str) -> Path:
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    return EXPORT_DIR / file_name


def _to_result(file_name: str) -> ExportResult:
    return ExportResult(file_name=file_name, download_url=f"/exports/{file_name}")


def _write_pdf(file_path: Path, package: dict) -> None:
    _register_pdf_fonts(package["language"])
    styles = getSampleStyleSheet()
    if package["language"] == "zh-CN":
        for style_name in ("Title", "Heading1", "Heading2", "Normal", "BodyText"):
            styles[style_name].fontName = "STSong-Light"

    story = [
        Paragraph(_document_title(package), styles["Title"]),
        Paragraph("Venture Copilot", styles["Heading2"]),
        Paragraph(_date_line(package), styles["Normal"]),
        Spacer(1, 18),
    ]

    for index, section in enumerate(package["business_plan_sections"], start=1):
        _pdf_section(
            story,
            styles,
            f"{index} {section['heading']}",
            [_export_section_content(section, package)],
        )
        if section["key"] == "financial_forecast":
            _pdf_finance_tables(story, package)
    if package.get("review_score"):
        _pdf_review_section(story, styles, package["review_score"], package["language"])

    SimpleDocTemplate(str(file_path), pagesize=letter).build(story)


def _pdf_section(story: list, styles, heading: str, paragraphs: list[str]) -> None:
    story.append(Paragraph(heading, styles["Heading1"]))
    for text in paragraphs:
        if text:
            story.append(Paragraph(str(text), styles["BodyText"]))
            story.append(Spacer(1, 8))


def _write_docx(file_path: Path, package: dict) -> None:
    document = Document()
    document.add_heading(_document_title(package), level=0)
    document.add_paragraph("Venture Copilot")
    document.add_paragraph(_date_line(package))

    for index, section in enumerate(package["business_plan_sections"], start=1):
        if section["key"] == "financial_forecast":
            _docx_financial_section(document, f"{index} {section['heading']}", section, package)
        else:
            _docx_section(
                document,
                f"{index} {section['heading']}",
                [_export_section_content(section, package)],
            )
    if package.get("review_score"):
        _docx_review_section(document, package["review_score"], package["language"])

    document.save(file_path)


def _pdf_review_section(story: list, styles, review: dict, language: str) -> None:
    headings = _review_headings(language)
    _pdf_section(story, styles, headings["score"], [f"{headings['overall']}：{review.get('overall_score')}"])
    _pdf_section(story, styles, headings["warnings"], review.get("warnings") or ["无"])
    _pdf_section(story, styles, headings["suggestions"], review.get("suggestions") or ["无"])


def _docx_review_section(document: Document, review: dict, language: str) -> None:
    headings = _review_headings(language)
    document.add_heading(headings["score"], level=1)
    document.add_paragraph(f"{headings['overall']}：{review.get('overall_score')}")
    document.add_heading(headings["warnings"], level=1)
    for item in review.get("warnings") or ["无"]:
        document.add_paragraph(str(item))
    document.add_heading(headings["suggestions"], level=1)
    for item in review.get("suggestions") or ["无"]:
        document.add_paragraph(str(item))


def _review_headings(language: str) -> dict:
    if language == "zh-CN":
        return {
            "score": "项目评分",
            "overall": "综合评分",
            "warnings": "风险提示",
            "suggestions": "改进建议",
        }
    return {
        "score": "Project Score",
        "overall": "Overall Score",
        "warnings": "Warnings",
        "suggestions": "Suggestions",
    }


def _docx_section(document: Document, heading: str, paragraphs: list[str]) -> None:
    document.add_heading(heading, level=1)
    for text in paragraphs:
        if text:
            document.add_paragraph(str(text))


def _docx_financial_section(document: Document, heading: str, section: dict, package: dict) -> None:
    document.add_heading(heading, level=1)
    for block in _export_section_content(section, package).split("\n\n"):
        if block.strip():
            document.add_paragraph(block.strip())
    tables = package.get("financial_narrative", {}).get("tables", {})
    _docx_table(document, "利润表", tables.get("income_statement") or [], ["month", "revenue", "gross_profit", "operating_profit", "net_profit"])
    _docx_table(document, "现金流量表", tables.get("cash_flow_statement") or [], ["month", "opening_cash", "operating_cash_flow", "financing_cash_flow", "ending_cash"])
    _docx_table(document, "资产负债表", tables.get("balance_sheet") or [], ["month", "cash", "assets", "liabilities", "equity"])
    _docx_table(document, "敏感性分析", tables.get("sensitivity_analysis") or [], ["scenario", "monthly_growth_rate", "total_revenue", "ending_cash", "funding_needed"])
    _docx_finance_interpretation(document, package.get("financial_interpretation") or {})
    _docx_chart_images(document, package.get("chart_files") or {})


def _docx_table(document: Document, title: str, rows: list[dict], columns: list[str]) -> None:
    if not rows:
        return
    document.add_heading(title, level=2)
    sample = rows[:12] if title != "敏感性分析" else rows
    table = document.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    headers = table.rows[0].cells
    for index, column in enumerate(columns):
        headers[index].text = _finance_label(column)
    for row in sample:
        cells = table.add_row().cells
        for index, column in enumerate(columns):
            cells[index].text = str(row.get(column, ""))
    if len(rows) > len(sample):
        document.add_paragraph(f"表格展示前 {len(sample)} 行，完整数据保存在系统财务预测 JSON 中。")


def _pdf_finance_tables(story: list, package: dict) -> None:
    tables = package.get("financial_narrative", {}).get("tables", {})
    for title, key, columns in [
        ("利润表", "income_statement", ["month", "revenue", "gross_profit", "operating_profit", "net_profit"]),
        ("现金流量表", "cash_flow_statement", ["month", "opening_cash", "operating_cash_flow", "ending_cash"]),
        ("资产负债表", "balance_sheet", ["month", "cash", "assets", "liabilities", "equity"]),
        ("敏感性分析", "sensitivity_analysis", ["scenario", "total_revenue", "ending_cash", "funding_needed"]),
    ]:
        rows = tables.get(key) or []
        if not rows:
            continue
        story.append(Paragraph(title, getSampleStyleSheet()["Heading2"]))
        sample = rows[:8] if key != "sensitivity_analysis" else rows
        data = [[_finance_label(column) for column in columns]]
        data.extend([[str(row.get(column, "")) for column in columns] for row in sample])
        story.append(Table(data, repeatRows=1))
        story.append(Spacer(1, 10))


def _join(items: list) -> str:
    return "; ".join(str(item) for item in items)


def _format_summary(summary: dict, language: str = "en-US") -> str:
    if language == "zh-CN":
        labels = {
            "total_revenue": "总收入",
            "total_net_profit": "净利润",
            "ending_cash": "期末现金",
            "breakeven_month": "盈亏平衡月份",
            "funding_needed": "资金缺口",
        }
        return "；".join(f"{labels.get(key, key)}：{value}" for key, value in summary.items())
    return "; ".join(f"{key}: {value}" for key, value in summary.items())


def _format_citations(citations: list[dict]) -> list[str]:
    if not citations:
        return ["暂无引用。"]
    return [
        f"{item.get('title')} — {item.get('source')} ({item.get('source_type')}): {item.get('url')}"
        for item in citations
    ]


def _normalize_business_plan_content(content: dict) -> dict:
    if isinstance(content.get("sections"), list):
        return {
            "language": content.get("language", "zh-CN"),
            "template_type": content.get("template_type", "generic_startup"),
            "selected_sections": content.get("selected_sections", []),
            "sections": content.get("sections", []),
        }

    sections = [
        {"key": key, "heading": key.replace("_", " ").title(), "content": value}
        for key, value in content.items()
    ]
    return {
        "language": "en-US",
        "template_type": "generic_startup",
        "selected_sections": [section["key"] for section in sections],
        "sections": sections,
    }


def _export_section_content(section: dict, package: dict) -> str:
    key = section["key"]
    language = package["language"]
    content = normalize_to_language(section["content"], language)

    if key == "financial_forecast":
        if language == "zh-CN":
            narrative = package.get("financial_narrative", {}).get("content") or content
            assumptions = _format_assumptions(package.get("financial_assumptions") or {})
            return normalize_to_language(f"财务假设\n{assumptions}\n\n{narrative}", language)
        return (
            f"{content}\n"
            f"Financial summary: breakeven month {package['breakeven']}; "
            f"funding needed {package['funding_needed']}; "
            f"{_format_summary(package['forecast_summary'], language)}"
        )

    if key == "market_analysis":
        if language == "zh-CN":
            return (
                f"{content}\n"
                f"研究补充：市场规模：{normalize_to_language(package['market_size'], language)}；"
                f"行业趋势：{_join(package['industry_trends'])}；"
                f"竞品：{_join(package['competitors'])}"
            )
        return (
            f"{content}\n"
            f"Research supplement: market size: {package['market_size']}; "
            f"industry trends: {_join(package['industry_trends'])}; "
            f"competitors: {_join(package['competitors'])}"
        )

    if key == "risk_analysis":
        if language == "zh-CN":
            return f"{content}\n研究风险：{_join(package['research_risks'])}"
        return f"{content}\nResearch risks: {_join(package['research_risks'])}"

    if key == "references":
        return content + "\n" + "\n".join(_format_citations(package["citations"]))

    return content


def _finance_label(key: str) -> str:
    return {
        "month": "月份",
        "revenue": "收入",
        "variable_costs": "变动成本",
        "fixed_costs": "固定成本",
        "gross_profit": "毛利润",
        "operating_profit": "营业利润",
        "net_profit": "净利润",
        "opening_cash": "期初现金",
        "operating_cash_flow": "经营现金流",
        "financing_cash_flow": "融资现金流",
        "ending_cash": "期末现金",
        "cash": "现金",
        "assets": "资产",
        "liabilities": "负债",
        "equity": "权益",
        "scenario": "情景",
        "monthly_growth_rate": "月增长率",
        "total_revenue": "总收入",
        "funding_needed": "资金需求",
    }.get(key, key)


def _format_assumptions(assumptions: dict) -> str:
    if not assumptions:
        return "缺少已确认财务假设。"
    labels = {
        "price": "售价",
        "unit_cost": "单位成本",
        "first_month_units": "预计首月销量",
        "growth_rate": "增长率",
        "startup_funds": "启动资金",
        "fixed_costs": "固定成本",
        "marketing_budget": "营销预算",
        "employee_costs": "员工成本",
        "forecast_months": "预测周期",
        "financing_amount": "融资金额",
        "r_and_d_costs": "研发成本",
        "repeat_purchase_rate": "复购率",
        "conversion_rate": "转化率",
        "ltv": "LTV",
        "cac": "CAC",
        "channel_costs": "渠道成本",
    }
    return "\n".join(f"{labels.get(key, key)}：{value}" for key, value in assumptions.items() if value not in (None, ""))


def _docx_finance_interpretation(document: Document, interpretation: dict) -> None:
    if not interpretation:
        return
    document.add_heading("财务解释与建议", level=2)
    for title, key in [("优势", "strengths"), ("弱点", "weaknesses"), ("风险", "risks"), ("建议", "recommendations")]:
        items = interpretation.get(key) or []
        if items:
            document.add_paragraph(title)
            for item in items:
                document.add_paragraph(str(item), style=None)


def _docx_chart_images(document: Document, chart_files: dict) -> None:
    if not chart_files:
        return
    captions = {
        "revenue_trend": "图：收入趋势",
        "cash_flow": "图：现金余额趋势",
        "cost_structure": "图：成本结构",
        "sensitivity": "图：敏感性分析",
    }
    document.add_heading("财务图表", level=2)
    for key, path in chart_files.items():
        try:
            document.add_picture(path, width=Inches(5.8))
            document.add_paragraph(captions.get(key, key))
        except Exception:
            document.add_paragraph(f"{captions.get(key, key)}：图表文件暂不可用。")


def _document_title(package: dict) -> str:
    if package["language"] == "zh-CN":
        return f"{package['project_title']} 创业计划书"
    return f"{package['project_title']} Business Plan"


def _date_line(package: dict) -> str:
    if package["language"] == "zh-CN":
        return f"生成日期：{package['generated_date']}"
    return f"Generated date: {package['generated_date']}"


def _register_pdf_fonts(language: str) -> None:
    if language == "zh-CN":
        try:
            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        except KeyError:
            pass
